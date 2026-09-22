"""CSV/XLSX contact importer used by Sekaro 0.3."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

MAX_IMPORT_BYTES = 15 * 1024 * 1024
MAX_IMPORT_ROWS = 100_000
PREVIEW_ROWS = 8

_EMAIL_ALIASES = {
    "email", "e-mail", "e mail", "mail", "email address", "email_address",
    "adres email", "adres e-mail", "adres e mail", "adres mail",
    "emailadresse", "e-mail-adresse", "mailadresse",
}
_NAME_ALIASES = {
    "name", "full name", "contact name", "contact_name",
    "imię i nazwisko", "imie i nazwisko", "nazwa", "kontakt",
    "ansprechpartner", "kontaktname",
}


@dataclass
class ParsedTable:
    headers: list[str]
    rows: list[dict[str, str]]
    filename: str
    sheet_name: str | None = None


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value).strip()


def _unique_headers(raw: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    headers: list[str] = []
    for idx, value in enumerate(raw, start=1):
        base = _stringify(value).strip() or f"column_{idx}"
        count = seen.get(base.lower(), 0) + 1
        seen[base.lower()] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    return headers


def _decode_csv(contents: bytes) -> str:
    last_error: Exception | None = None
    for enc in ("utf-8-sig", "cp1250", "windows-1252", "latin-1"):
        try:
            return contents.decode(enc)
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Could not decode text file: {last_error}")


def _parse_csv(filename: str, contents: bytes) -> ParsedTable:
    text = _decode_csv(contents)
    if not text.strip():
        raise ValueError("File is empty")

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        if "\t" in sample and "," not in sample:
            dialect = csv.excel_tab

    reader = csv.reader(io.StringIO(text), dialect=dialect)
    try:
        raw_headers = next(reader)
    except StopIteration:
        raise ValueError("File is empty")

    headers = _unique_headers(raw_headers)
    rows: list[dict[str, str]] = []
    for values in reader:
        if len(rows) >= MAX_IMPORT_ROWS:
            raise ValueError(f"Import exceeds the limit of {MAX_IMPORT_ROWS} rows")
        padded = list(values[: len(headers)]) + [""] * max(0, len(headers) - len(values))
        row = {headers[i]: _stringify(padded[i]) for i in range(len(headers))}
        if any(v.strip() for v in row.values()):
            rows.append(row)

    return ParsedTable(headers=headers, rows=rows, filename=filename)


def _parse_xlsx(filename: str, contents: bytes) -> ParsedTable:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("XLSX support is not installed") from exc

    try:
        wb = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read Excel file: {exc}") from exc

    if not wb.sheetnames:
        raise ValueError("Excel workbook has no worksheets")

    ws = wb[wb.sheetnames[0]]
    iterator = ws.iter_rows(values_only=True)
    try:
        raw_headers = list(next(iterator))
    except StopIteration:
        raise ValueError("Excel worksheet is empty")

    headers = _unique_headers(raw_headers)
    rows: list[dict[str, str]] = []
    for values in iterator:
        if len(rows) >= MAX_IMPORT_ROWS:
            raise ValueError(f"Import exceeds the limit of {MAX_IMPORT_ROWS} rows")
        values = list(values)
        padded = values[: len(headers)] + [None] * max(0, len(headers) - len(values))
        row = {headers[i]: _stringify(padded[i]) for i in range(len(headers))}
        if any(v.strip() for v in row.values()):
            rows.append(row)

    return ParsedTable(
        headers=headers,
        rows=rows,
        filename=filename,
        sheet_name=ws.title,
    )


def parse_contact_file(filename: str, contents: bytes) -> ParsedTable:
    if len(contents) > MAX_IMPORT_BYTES:
        raise ValueError(
            f"File is too large. Maximum size is {MAX_IMPORT_BYTES // (1024 * 1024)} MB"
        )

    suffix = Path(filename or "").suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx(filename, contents)
    if suffix in {".csv", ".tsv", ".txt", ""}:
        return _parse_csv(filename, contents)
    if suffix == ".xls":
        raise ValueError("Old .xls files are not supported. Save the file as .xlsx or CSV.")
    raise ValueError("Supported formats: .xlsx, .xlsm, .csv, .tsv, .txt")


def _norm_header(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().lower())


def slugify_custom_field(header: str) -> str:
    value = header.strip().lower()
    value = re.sub(r"[^0-9a-zA-ZąćęłńóśźżĄĆĘŁŃÓŚŹŻ]+", "_", value)
    value = value.strip("_")
    return value[:64] or "field"


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    """Return source-column to target mapping.

    Targets are email, name, custom:<key>, or skip.
    Every non-core column is preserved as a custom field by default.
    """
    result: dict[str, str] = {}
    email_taken = False
    name_taken = False

    for header in headers:
        normalized = _norm_header(header)
        if not email_taken and normalized in _EMAIL_ALIASES:
            result[header] = "email"
            email_taken = True
        elif not name_taken and normalized in _NAME_ALIASES:
            result[header] = "name"
            name_taken = True
        else:
            result[header] = f"custom:{slugify_custom_field(header)}"
    return result


def apply_mapping(
    row: dict[str, str],
    mapping: dict[str, str],
) -> tuple[str, str, dict[str, str]]:
    email = ""
    name = ""
    custom: dict[str, str] = {}

    for source, target in mapping.items():
        value = (row.get(source) or "").strip()
        if not target or target == "skip":
            continue
        if target == "email":
            email = value.lower()
        elif target == "name":
            name = value
        elif target.startswith("custom:"):
            key = target.split(":", 1)[1].strip()
            if key and value:
                custom[key] = value

    return email, name, custom


def validate_mapping(headers: list[str], mapping: dict[str, str]) -> None:
    unknown = [source for source in mapping if source not in headers]
    if unknown:
        raise ValueError(f"Mapping contains unknown columns: {', '.join(unknown)}")

    email_sources = [source for source, target in mapping.items() if target == "email"]
    if len(email_sources) != 1:
        raise ValueError("Exactly one source column must be mapped to email")

    name_sources = [source for source, target in mapping.items() if target == "name"]
    if len(name_sources) > 1:
        raise ValueError("At most one source column can be mapped to name")

    custom_keys: list[str] = []
    for target in mapping.values():
        if target.startswith("custom:"):
            key = target.split(":", 1)[1].strip()
            if not key:
                raise ValueError("Custom field key cannot be empty")
            custom_keys.append(key.lower())
    if len(custom_keys) != len(set(custom_keys)):
        raise ValueError("Custom field keys must be unique")
